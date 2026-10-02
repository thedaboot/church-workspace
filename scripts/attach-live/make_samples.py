# 첨부 실측용 샘플 — 레포 밖 폴더에 만든다(파일을 레포에 두지 않는다 · HANDOFF §1-9).
#   python scripts/attach-live/make_samples.py <출력 폴더>      (ffmpeg · python-docx · openpyxl · python-pptx 필요)
# 옛 바이너리(.doc·.xls·.ppt)와 heic은 만들 도구가 없어 빠진다. .hwp는 머리말만 흉내 낸 가짜다.
import os, sys, struct, zlib, json, random, subprocess
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as PI

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)
P = lambda n: os.path.join(OUT, n)

def png_bytes(w, h, seed=1, noise=False):
    rnd = random.Random(seed)
    rows = []
    for y in range(h):
        row = bytearray([0])
        for x in range(w):
            if noise:
                row += bytes([rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)])
            else:
                row += bytes([(x * 255) // max(1, w - 1), (y * 255) // max(1, h - 1), 160])
        rows.append(bytes(row))
    raw = b''.join(rows)
    def chunk(t, d): return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw, 9)) + chunk(b'IEND', b'')

open(P('검증 사진.png'), 'wb').write(png_bytes(640, 420))
ff = lambda *a: subprocess.run(['ffmpeg', '-y', '-hide_banner', '-loglevel', 'error', *a], check=True)
ff('-i', P('검증 사진.png'), P('sample.jpg'))
ff('-i', P('검증 사진.png'), P('sample.webp'))
ff('-i', P('검증 사진.png'), P('sample.gif'))
ff('-i', P('검증 사진.png'), P('sample.bmp'))
try: ff('-i', P('검증 사진.png'), '-c:v', 'libaom-av1', '-still-picture', '1', P('sample.avif'))
except Exception as e: print('avif 실패', e)
open(P('sample.svg'), 'w', encoding='utf8').write('<svg xmlns="http://www.w3.org/2000/svg" width="320" height="200"><rect width="320" height="200" fill="#3f6fc4"/><text x="20" y="110" font-size="28" fill="#fff">SVG 검증</text></svg>')

# 영상·소리
ff('-f', 'lavfi', '-i', 'testsrc=size=320x240:rate=15', '-f', 'lavfi', '-i', 'sine=frequency=440', '-t', '3', '-c:v', 'libx264', '-pix_fmt', 'yuv420p', '-c:a', 'aac', '-shortest', P('sample.mp4'))
ff('-f', 'lavfi', '-i', 'testsrc=size=320x240:rate=15', '-t', '2', '-c:v', 'libx264', '-pix_fmt', 'yuv420p', P('sample.mov'))
ff('-f', 'lavfi', '-i', 'testsrc=size=320x240:rate=15', '-t', '2', '-c:v', 'libvpx', P('sample.webm'))
ff('-f', 'lavfi', '-i', 'sine=frequency=440', '-t', '2', '-c:a', 'libmp3lame', P('sample.mp3'))
ff('-f', 'lavfi', '-i', 'sine=frequency=440', '-t', '2', '-c:a', 'aac', P('sample.m4a'))
ff('-f', 'lavfi', '-i', 'sine=frequency=440', '-t', '2', P('sample.wav'))
ff('-f', 'lavfi', '-i', 'sine=frequency=440', '-t', '2', '-c:a', 'libvorbis', P('sample.ogg'))
# 3MB 넘는 영상 — 큰 파일 갈래(Storage 거쳐 uploadFromUrl)
ff('-f', 'lavfi', '-i', 'nullsrc=size=640x480:rate=30,geq=random(1)*255:128:128', '-t', '4', '-c:v', 'libx264', '-preset', 'ultrafast', '-crf', '30', '-pix_fmt', 'yuv420p', P('big-video.mp4'))

# 글자류
open(P('sample.txt'), 'w', encoding='utf8').write('첨부 검증 텍스트 TXT-MARK\n둘째 줄\n')
open(P('sample.md'), 'w', encoding='utf8').write('# 마크다운 검증\n\n**MD-MARK** 굵게\n\n- 항목 하나\n')
open(P('sample.csv'), 'w', encoding='utf8').write('이름,값\n검증,CSV-MARK\n둘,2\n')
open(P('sample.tsv'), 'w', encoding='utf8').write('이름\t값\n검증\tTSV-MARK\n')
open(P('sample.json'), 'w', encoding='utf8').write(json.dumps({'mark': 'JSON-MARK', '값': 1}, ensure_ascii=False))
open(P('sample.html'), 'w', encoding='utf8').write('<!doctype html><html><body><h1 id="m">HTML-MARK</h1><script>document.body.dataset.ran="1"</script></body></html>')

# PDF(손으로 — 두 쪽)
def pdf(path, pages):
    objs = []
    objs.append('<< /Type /Catalog /Pages 2 0 R >>')
    kids = ' '.join(f'{3 + i * 2} 0 R' for i in range(pages))
    objs.append(f'<< /Type /Pages /Kids [{kids}] /Count {pages} >>')
    font_id = 3 + pages * 2
    for i in range(pages):
        content = f'BT /F1 36 Tf 72 700 Td (PDF-MARK page {i + 1}) Tj ET'
        objs.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + i * 2} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>')
        objs.append(f'<< /Length {len(content)} >>\nstream\n{content}\nendstream')
    objs.append('<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>')
    out = b'%PDF-1.4\n'
    offs = []
    for n, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f'{n} 0 obj\n{o}\nendobj\n'.encode()
    x = len(out)
    out += f'xref\n0 {len(objs) + 1}\n0000000000 65535 f \n'.encode()
    for o in offs: out += f'{o:010d} 00000 n \n'.encode()
    out += f'trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n'.encode()
    open(path, 'wb').write(out)
pdf(P('검증 문서.pdf'), 2)
pdf(P('songform.pdf'), 1)

# 오피스
d = Document(); d.add_heading('워드 검증 DOCX-MARK', 1); d.add_paragraph('본문 문단입니다.'); d.save(P('검증 워드.docx'))
wb = Workbook(); wb.active.append(['매크로', 'XLSM-MARK']); wb.save(P('검증 매크로.xlsm'))
wb = Workbook(); ws = wb.active; ws.title = '첫 시트'; ws.append(['이름', '값']); ws.append(['검증', 'XLSX-MARK']); wb.create_sheet('둘째 시트').append(['둘']); wb.save(P('검증 엑셀.xlsx'))
pr = Presentation(); s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = 'PPT 검증 PPTX-MARK'; s.placeholders[1].text = '첫 장'
s2 = pr.slides.add_slide(pr.slide_layouts[1]); s2.shapes.title.text = '둘째 장'; pr.save(P('검증 슬라이드.pptx'))
# 3MB 넘는 워드(그림 포함) — 큰 파일 갈래 + 변환 사본
open(P('noise.png'), 'wb').write(png_bytes(1100, 1100, seed=7, noise=True))
d = Document(); d.add_heading('큰 워드 BIGDOCX-MARK', 1); d.add_picture(P('noise.png'), width=Inches(5)); d.save(P('큰 워드.docx'))
os.remove(P('noise.png'))
# 큐시트(주보)
d = Document(); d.add_heading('2099-12-27 큐시트 CUE-MARK', 1); t = d.add_table(rows=2, cols=3); t.cell(0, 0).text = '순서'; t.cell(0, 1).text = '내용'; t.cell(0, 2).text = '담당'; d.save(P('20991227_검증 큐시트.docx'))
# 그 밖 — 미리보기 없는 형식
import zipfile
with zipfile.ZipFile(P('sample.zip'), 'w') as z: z.writestr('a.txt', 'zip 안')
open(P('sample.hwp'), 'wb').write(b'HWP Document File' + b'\0' * 64)   # 진짜 hwp 아님(헤더 흉내) — 업로드·폴더·삭제만 본다
for f in sorted(os.listdir(OUT)): print(f, os.path.getsize(P(f)))
